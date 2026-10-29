"""Excel 解析：每个 sheet 一个章节，每 20 行一块并重复表头。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from kbqa.core.ingest.cleaner import clean_text
from kbqa.core.ingest.types import Block, EmptyDocumentError, ParsedDocument, ParseError

ROWS_PER_BLOCK = 20


def _to_md(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    head, body = rows[0], rows[1:]
    lines = ["| " + " | ".join(head) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)


class XlsxParser:
    extensions = ("xlsx",)

    def parse(self, path: Path) -> ParsedDocument:
        try:
            wb = load_workbook(str(path), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise ParseError(f"xlsx 打开失败：{exc}") from exc

        blocks: list[Block] = []
        parts: list[str] = []
        cursor = 0
        idx = 0

        for ws in wb.worksheets:
            rows = [[("" if c is None else str(c).strip()) for c in row]
                    for row in ws.iter_rows(values_only=True)]
            rows = [r for r in rows if any(r)]
            if not rows:
                continue
            header, body = rows[0], rows[1:]
            for i in range(0, max(len(body), 1), ROWS_PER_BLOCK):
                part = body[i:i + ROWS_PER_BLOCK]
                text = clean_text(_to_md([header] + part)) if part else clean_text(_to_md([header]))
                if not text:
                    continue
                if parts:
                    cursor += 2
                blocks.append(Block("table", text, cursor, cursor + len(text),
                                    para_index=idx, meta={"sheet": ws.title}))
                parts.append(text)
                cursor += len(text)
                idx += 1

        wb.close()
        plain = "\n\n".join(parts)
        if len(plain.strip()) < 20:
            raise EmptyDocumentError(f"表格内容为空：{path.name}")
        return ParsedDocument(blocks=blocks, plain_text=plain, doc_meta={"tables": idx})
